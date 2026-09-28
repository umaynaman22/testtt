import io
import re
import time
from datetime import date
from pathlib import Path

import pytest

from rental_tracker import db as dbmod
from rental_tracker.demo import build_demo
from rental_tracker.web import create_app

TODAY = date(2026, 9, 27)
TOKEN = "test-token"


@pytest.fixture
def app(data_dir):
    build_demo(data_dir, today=TODAY)
    return create_app(data_dir, launch_token=TOKEN, today_fn=lambda: TODAY, testing=True)


@pytest.fixture
def client(app):
    c = app.test_client()
    assert c.get(f"/auth?token={TOKEN}").status_code == 302
    return c


def csrf(client) -> str:
    html = client.get("/").get_data(as_text=True)
    return re.search(r'"X-CSRF-Token": "([^"]+)"', html).group(1)


def q(app, sql):
    conn = dbmod.connect(app.extensions["rental_tracker"].data.db)
    try:
        return [r[0] for r in conn.execute(sql)]
    finally:
        conn.close()


def test_requires_launch_token(app):
    c = app.test_client()
    assert c.get("/").status_code == 403
    assert c.get("/auth?token=wrong").status_code == 403
    assert c.get(f"/auth?token={TOKEN}").status_code == 302
    assert c.get("/").status_code == 200


def test_rejects_foreign_host(client):
    assert client.get("/", headers={"Host": "evil.example.com"}).status_code == 400
    assert client.get("/", headers={"Host": "127.0.0.1:5000"}).status_code != 400  # allowed host


def test_post_requires_csrf(client):
    assert client.post("/properties/new", data={"name": "x"}).status_code == 400
    assert client.post("/properties/new", data={"name": "x", "csrf_token": csrf(client)}).status_code == 302


def test_every_page_renders(app, client):
    pid, unit, lease, tid, pay = (q(app, f"SELECT id FROM {t} LIMIT 1")[0]
                                  for t in ("properties", "units", "leases", "tenants", "payments"))
    pages = ["/", "/properties", "/properties?sort=balance", "/properties/new", f"/properties/{pid}",
             f"/properties/{pid}/edit", "/tenants", "/tenants?q=lee",
             "/tenants?format=xlsx", "/tenants/new", f"/tenants/new?unit_id={unit}", f"/tenants/{tid}",
             f"/tenants/{tid}/edit", "/rent-day", "/rent-day?period=2026-08&show=unpaid", "/late", "/payments",
             "/payments?format=xlsx", "/reports/everything.xlsx", f"/payments/{pay}/receipt", "/reports", "/settings", "/search?q=rizal",
             "/search?q=zzzz"]
    for key in ("rent-roll", "aging", "collections"):
        pages += [f"/reports/{key}", f"/reports/{key}?format=xlsx", f"/reports/{key}?property={pid}"]
    for lid in q(app, "SELECT id FROM leases"):
        pages += [f"/leases/{lid}", f"/leases/{lid}/edit", f"/leases/{lid}/statement"]
    for page in pages:
        resp = client.get(page)
        assert resp.status_code in (200, 302), f"{page} -> {resp.status_code}"
    assert client.get("/properties/999999").status_code == 404
    assert client.get("/leases/999999").status_code == 404
    assert client.get("/reports/deposits").status_code == 404
    assert client.get("/reports/vacancy").status_code == 404
    for gone in ("/expenses", "/vendors", "/owners", "/import", "/backups", "/audit", "/leases"):
        assert client.get(gone).status_code == 404, gone


def test_no_required_fields_anywhere():
    root = Path(__file__).resolve().parents[2] / "src" / "rental_tracker" / "web" / "templates"
    for path in root.rglob("*.html"):
        text = path.read_text(encoding="utf-8")
        assert not re.search(r"<(input|select|textarea)[^>]*\srequired\b", text), path
        assert "<em>*</em>" not in text, path


def test_no_external_urls_in_templates_or_static():
    root = Path(__file__).resolve().parents[2] / "src" / "rental_tracker" / "web"
    for path in list((root / "templates").rglob("*.html")) + [root / "static" / "app.css", root / "static" / "app.js"]:
        text = path.read_text(encoding="utf-8")
        assert not re.search(r"(src|href)=[\"']https?://", text), path


def test_simple_workflow(app, client):
    token = csrf(client)
    # a property with nothing filled in
    r = client.post("/properties/new", data={"csrf_token": token})
    assert r.status_code == 302
    pid = int(r.headers["Location"].rstrip("/").split("/")[-1])
    unit = q(app, f"SELECT id FROM units WHERE property_id = {pid}")[0]
    assert "Add tenant" in client.get(f"/properties/{pid}").get_data(as_text=True)
    # a tenant who has lived there since last year: rent is billed from the move-in date, plus a debt
    r = client.post("/tenants/new", data={"csrf_token": token, "unit_id": unit, "name": "Rita Moreno",
                                          "rent": "1,200", "moved_in": "2025-05-01", "debt": "300"})
    assert r.status_code == 302, r.get_data(as_text=True)[:2000]
    lid = int(r.headers["Location"].rstrip("/").split("/")[-1])
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert "Rita Moreno" in page and ">Debt<" in page and "Fill in past rent" in page
    assert q(app, f"SELECT COUNT(*) FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'")[0] == 17  # May 2025-Sep 2026
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0] == 17 * 120000 + 30000
    # they had paid all that rent: fill in the payment history up to today
    r = client.post(f"/leases/{lid}/fill-rent", data={"csrf_token": token, "method": "cash"}, follow_redirects=True)
    assert "Marked 17 rent bills as paid" in r.get_data(as_text=True)
    assert q(app, f"SELECT COUNT(*) FROM payments WHERE lease_id = {lid} AND method = 'cash'")[0] == 17
    assert q(app, f"SELECT MIN(received_date) FROM payments WHERE lease_id = {lid}") == ["2025-05-01"]
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0] == 30000  # the debt is left
    assert "Fill in past rent" not in client.get(f"/leases/{lid}").get_data(as_text=True)
    # a payment with only an amount
    client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": "300"})
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0] == 0
    assert q(app, f"SELECT method FROM payments WHERE lease_id = {lid} ORDER BY id DESC LIMIT 1")[0] == "other"
    # Collect rent quick entry, and the typo guard
    r = client.post("/rent-day/pay", data={"lease_id": lid, "amount": "310000", "period": "2026-09"},
                    headers={"HX-Request": "true", "X-CSRF-Token": token})
    assert "far more than this lease owes" in r.get_data(as_text=True)
    # edit: change rent after billing -> starts with the next bill
    client.post(f"/leases/{lid}/edit", data={"csrf_token": token, "name": "Rita M", "rent": "1300"})
    assert q(app, f"SELECT rent_cents FROM lease_rent_changes WHERE lease_id = {lid}") == [130000]
    assert q(app, "SELECT first_name FROM tenants WHERE last_name = 'M'") == ["Rita"]
    # add a debt with a note, then one with only an amount
    client.post(f"/leases/{lid}/debt", data={"csrf_token": token, "amount": "75", "note": "broken window"})
    client.post(f"/leases/{lid}/debt", data={"csrf_token": token, "amount": "25"})
    assert q(app, f"SELECT description FROM charges WHERE lease_id = {lid} AND charge_type = 'other' ORDER BY id") == \
        ["Debt — broken window", "Debt"]
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0] == 10000
    assert "Debt — broken window" in client.get(f"/leases/{lid}").get_data(as_text=True)
    client.post(f"/leases/{lid}/debt", data={"csrf_token": token})  # nothing entered: nothing added
    assert q(app, f"SELECT COUNT(*) FROM charges WHERE lease_id = {lid} AND charge_type = 'other'")[0] == 2


def test_removed_fields_and_actions(app, client):
    lid = q(app, "SELECT id FROM leases WHERE status = 'active' LIMIT 1")[0]
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert "Add debt" in page
    for gone in ("Add a charge", "Give a credit", "Change the rent", "Security deposit", "Deposit held",
                 "Moved out", "mailto:"):
        assert gone not in page, gone
    for url in (f"/leases/{lid}/charge", f"/leases/{lid}/credit", f"/leases/{lid}/deposit",
                f"/leases/{lid}/rent-change", f"/leases/{lid}/move-out"):
        assert client.post(url, data={"csrf_token": csrf(client)}).status_code == 404, url
    forms = client.get("/tenants/new").get_data(as_text=True) + client.get(f"/leases/{lid}/edit").get_data(as_text=True)
    for gone in ('name="email"', 'name="second_name"', 'name="end_date"', 'name="deposit"', "Already owes"):
        assert gone not in forms, gone
    assert 'name="debt"' in forms
    tid = q(app, "SELECT id FROM tenants LIMIT 1")[0]
    assert 'name="email"' not in client.get(f"/tenants/{tid}/edit").get_data(as_text=True)
    assert 'class="tabs"' not in client.get("/tenants").get_data(as_text=True)


def test_fill_in_past_rent_up_to_a_date_and_rebill_on_edit(app, client):
    token = csrf(client)
    pid = client.post("/properties/new", data={"csrf_token": token, "name": "5 Luna St"}).headers["Location"].split("/")[-1]
    unit = q(app, f"SELECT id FROM units WHERE property_id = {pid}")[0]
    lid = client.post("/tenants/new", data={"csrf_token": token, "unit_id": unit, "name": "Ana Cruz", "rent": "10000",
                                            "moved_in": "2026-01-01", "late_fee": "500"}).headers["Location"].split("/")[-1]
    assert q(app, f"SELECT COUNT(*) FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'")[0] == 9  # Jan-Sep
    # paid through June only
    client.post(f"/leases/{lid}/fill-rent", data={"csrf_token": token, "through": "2026-06-30", "method": "other",
                                                  "method_other": "GCash padala"})
    assert q(app, f"SELECT COUNT(*) FROM payments WHERE lease_id = {lid} AND method_other = 'GCash padala'")[0] == 6
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0] == 3 * 1000000
    late = client.get("/late").get_data(as_text=True)
    assert "Ana Cruz" in late  # July to September are still owed
    from rental_tracker.services import late_fees
    conn = dbmod.connect(app.extensions["rental_tracker"].data.db)
    periods = {c.period for c in late_fees.find_candidates(conn, TODAY) if c.lease_id == int(lid)}
    conn.close()
    assert periods == {"2026-07", "2026-08", "2026-09"}  # no late fees for the months filled in as paid
    r = client.post(f"/leases/{lid}/fill-rent", data={"csrf_token": token, "through": "2026-06-30"}, follow_redirects=True)
    assert "no unpaid rent up to that date" in r.get_data(as_text=True)
    # moving the move-in date later re-bills from then; payments stay
    client.post(f"/leases/{lid}/edit", data={"csrf_token": token, "name": "Ana Cruz", "rent": "10000",
                                             "moved_in": "2026-03-01"})
    assert q(app, f"SELECT MIN(period) FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'") == ["2026-03"]
    assert q(app, f"SELECT COUNT(*) FROM payments WHERE lease_id = {lid}")[0] == 6
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0] == 7 * 1000000 - 6 * 1000000


def _old_style_tenant(app, name):
    """A tenant added by an older version: moved in 2025-05-01, rent billed only from Sep 2026."""
    from rental_tracker.services import leases, portfolio, tenants
    conn = dbmod.connect(app.extensions["rental_tracker"].data.db)
    pid = portfolio.save_property(conn, None, {"name": f"{name}'s unit"})
    unit = conn.execute("SELECT id FROM units WHERE property_id = ?", (pid,)).fetchone()[0]
    tid = tenants.save_tenant(conn, None, {"first_name": name})
    lid = leases.create_lease(conn, unit_id=unit, tenants=[(tid, "primary")], start="2025-05-01", end=None,
                              rent_cents=500000, today=TODAY, billing_start="2026-09-01", move_in_date="2025-05-01")
    conn.close()
    return lid


def test_older_tenants_can_be_billed_from_move_in(app, client):
    """Tenants added before rent was billed from the move-in date can get their past months billed."""
    token = csrf(client)
    lid = _old_style_tenant(app, "Old")
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert "Fill in past rent" in page and "Bill rent since May 1, 2025" in page and "Mark rent as paid" not in page
    r = client.post(f"/leases/{lid}/bill-from-move-in", data={"csrf_token": token}, follow_redirects=True)
    page = r.get_data(as_text=True)
    assert "Rent is now billed from the move-in date" in page
    assert "Bill rent since" not in page and "Mark rent as paid" in page and "17 rent bills unpaid" in page
    assert q(app, f"SELECT MIN(period) FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'") == ["2025-05"]
    client.post(f"/leases/{lid}/fill-rent", data={"csrf_token": token, "method": "cash"})
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}") == [0]
    assert "Fill in past rent" not in client.get(f"/leases/{lid}").get_data(as_text=True)


def test_older_tenants_edit_checkbox(app, client):
    token = csrf(client)
    lid = _old_style_tenant(app, "Older")
    assert 'name="bill_from_move_in"' in client.get(f"/leases/{lid}/edit").get_data(as_text=True)
    client.post(f"/leases/{lid}/edit", data={"csrf_token": token, "moved_in": "2025-05-01"})  # saving alone changes nothing
    assert q(app, f"SELECT MIN(period) FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'") == ["2026-09"]
    client.post(f"/leases/{lid}/edit", data={"csrf_token": token, "moved_in": "2025-05-01", "bill_from_move_in": "1"})
    assert q(app, f"SELECT MIN(period) FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'") == ["2025-05"]
    assert q(app, f"SELECT billing_start_date FROM leases WHERE id = {lid}") == [None]
    assert 'name="bill_from_move_in"' not in client.get(f"/leases/{lid}/edit").get_data(as_text=True)


def test_mark_a_history_line_paid(app, client):
    token = csrf(client)
    pid = client.post("/properties/new", data={"csrf_token": token, "name": "8 Tala St"}).headers["Location"].split("/")[-1]
    unit = q(app, f"SELECT id FROM units WHERE property_id = {pid}")[0]
    lid = client.post("/tenants/new", data={"csrf_token": token, "unit_id": unit, "name": "Joy Bautista",
                                            "rent": "6000", "moved_in": "2026-07-01"}).headers["Location"].split("/")[-1]
    rent = dict(dbmod.connect(app.extensions["rental_tracker"].data.db).execute(
        f"SELECT period, id FROM charges WHERE lease_id = {lid} AND charge_type = 'rent'").fetchall())
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert page.count("Mark paid") == 3
    # August is paid, even though July is still open
    r = client.post(f"/leases/{lid}/pay-line/{rent['2026-08']}", data={"csrf_token": token, "amount": "6000"},
                    follow_redirects=True)
    page = r.get_data(as_text=True)
    assert "Rent 2026-08: paid." in page and page.count("Mark paid") == 2 and ">Paid</span>" in page
    assert q(app, f"SELECT received_date FROM payments WHERE lease_id = {lid}") == ["2026-08-01"]
    # part of September
    r = client.post(f"/leases/{lid}/pay-line/{rent['2026-09']}", data={"csrf_token": token, "amount": "2,500"},
                    follow_redirects=True)
    page = r.get_data(as_text=True)
    assert "Rent 2026-09: ₱2,500.00 paid, ₱3,500.00 left." in page and "₱2,500.00 paid</div>" in page
    assert 'value="3500.00"' in page  # the button now offers what's left
    # July is still the one that's owed (and late)
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}") == [600000 + 350000]
    r = client.post(f"/leases/{lid}/pay-line/{rent['2026-08']}", data={"csrf_token": token}, follow_redirects=True)
    assert "already paid" in r.get_data(as_text=True)
    other = q(app, f"SELECT id FROM charges WHERE lease_id <> {lid} LIMIT 1")[0]
    r = client.post(f"/leases/{lid}/pay-line/{other}", data={"csrf_token": token}, follow_redirects=True)
    assert q(app, f"SELECT COUNT(*) FROM payments WHERE lease_id = {lid}") == [2]  # someone else's line: nothing


def test_change_dates_in_payment_history(app, client):
    token = csrf(client)
    pay, lid = dbmod.connect(app.extensions["rental_tracker"].data.db).execute(
        "SELECT id, lease_id FROM payments ORDER BY id LIMIT 1").fetchone()
    charge = q(app, f"SELECT id FROM charges WHERE lease_id = {lid} AND charge_type = 'rent' ORDER BY due_date LIMIT 1")[0]
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert f"/leases/{lid}/line/payment/{pay}" in page and f"/leases/{lid}/line/charge/{charge}" in page
    r = client.post(f"/leases/{lid}/line/payment/{pay}", data={"csrf_token": token, "date": "2026-03-09"},
                    follow_redirects=True)
    assert "Saved" in r.get_data(as_text=True)
    assert q(app, f"SELECT received_date FROM payments WHERE id = {pay}") == ["2026-03-09"]
    client.post(f"/leases/{lid}/line/charge/{charge}", data={"csrf_token": token, "date": "2026-03-05"})
    assert q(app, f"SELECT due_date FROM charges WHERE id = {charge}") == ["2026-03-05"]
    r = client.post(f"/leases/{lid}/line/charge/{charge}", data={"csrf_token": token, "date": "31/31/2026"},
                    follow_redirects=True)
    assert "error" in r.get_data(as_text=True) and q(app, f"SELECT due_date FROM charges WHERE id = {charge}") == ["2026-03-05"]
    other = q(app, f"SELECT id FROM payments WHERE lease_id <> {lid} LIMIT 1")[0]
    client.post(f"/leases/{lid}/line/payment/{other}", data={"csrf_token": token, "date": "2020-01-01"})
    assert q(app, f"SELECT received_date FROM payments WHERE id = {other}") != ["2020-01-01"]  # not this tenant's
    assert client.post(f"/leases/{lid}/line/tenant/1", data={"csrf_token": token, "date": "2026-01-01"}).status_code == 404


def test_every_box_in_payment_history_can_be_changed(app, client):
    token = csrf(client)
    pid = client.post("/properties/new", data={"csrf_token": token, "name": "3 Dahlia St"}).headers["Location"].split("/")[-1]
    unit = q(app, f"SELECT id FROM units WHERE property_id = {pid}")[0]
    lid = client.post("/tenants/new", data={"csrf_token": token, "unit_id": unit, "name": "Nina Cruz", "rent": "9000",
                                            "moved_in": "2026-09-01"}).headers["Location"].split("/")[-1]
    client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": "9000", "method": "cash"})
    rent = q(app, f"SELECT id FROM charges WHERE lease_id = {lid}")[0]
    pay = q(app, f"SELECT id FROM payments WHERE lease_id = {lid}")[0]
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert page.count(f"/leases/{lid}/line/charge/{rent}") == 3  # date, what, charged
    assert page.count(f"/leases/{lid}/line/payment/{pay}") == 3  # date, how they paid, paid
    line = f"/leases/{lid}/line/charge/{rent}"
    client.post(line, data={"csrf_token": token, "description": "Rent for September (discounted)"})
    client.post(line, data={"csrf_token": token, "amount": "8,500"})
    assert q(app, f"SELECT description || '|' || amount_cents FROM charges WHERE id = {rent}") == [
        "Rent for September (discounted)|850000"]
    pline = f"/leases/{lid}/line/payment/{pay}"
    client.post(pline, data={"csrf_token": token, "amount": "8500"})
    client.post(pline, data={"csrf_token": token, "method": "other", "method_other": "Maya"})
    assert [tuple(r) for r in dbmod.connect(app.extensions["rental_tracker"].data.db).execute(
        f"SELECT amount_cents, method, method_other FROM payments WHERE id = {pay}")] == [(850000, "other", "Maya")]
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    assert "Rent for September (discounted)" in page and "Payment — Maya" in page
    assert q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}") == [0]  # balance follows
    r = client.post(pline, data={"csrf_token": token, "amount": "0"}, follow_redirects=True)
    assert "greater than zero" in r.get_data(as_text=True)
    client.post(pline, data={"csrf_token": token, "method": "gcash", "method_other": "ignored"})
    assert q(app, f"SELECT method || '/' || COALESCE(method_other, '') FROM payments WHERE id = {pay}") == ["gcash/"]
    client.post(line, data={"csrf_token": token, "description": ""})  # blank: back to the standard name
    assert q(app, f"SELECT description FROM charges WHERE id = {rent}") == [None]
    assert "rent" in client.get(f"/leases/{lid}").get_data(as_text=True).lower()


def test_review_fixes(app, client):
    """Bugs found in the code review stay fixed."""
    token = csrf(client)
    lid = q(app, "SELECT id FROM leases WHERE status = 'active' LIMIT 1")[0]
    big = "99999999999999999999"
    # numbers too big for the database: a message or an error page, never a crash
    r = client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": big}, follow_redirects=True)
    assert r.status_code == 200 and "enter an amount like" in r.get_data(as_text=True)
    for url in (f"/leases/{big}", f"/tenants/new?unit_id={big}&tenant_id={big}", f"/payments?property={big}"):
        assert client.get(url).status_code < 500, url
    # a unit's name and the "back" link are shown as text, not HTML
    pid = client.post("/properties/new", data={"csrf_token": token, "name": "Tom <b>Bold</b>"}).headers["Location"].split("/")[-1]
    unit = q(app, f"SELECT id FROM units WHERE property_id = {pid}")[0]
    new = client.post("/tenants/new", data={"csrf_token": token, "unit_id": unit, "name": "X Y"}).headers["Location"]
    assert "<b>Bold</b>" not in client.get(new).get_data(as_text=True)
    tid = q(app, "SELECT id FROM tenants LIMIT 1")[0]
    page = client.get(f'/tenants/{tid}/edit?back="><i>x</i>').get_data(as_text=True)
    assert "<i>x</i>" not in page
    assert 'href="//evil.example"' not in client.get(f"/tenants/{tid}/edit?back=//evil.example").get_data(as_text=True)
    # clearing a name gives "Unnamed tenant" rather than a blank
    client.post(f"/leases/{lid}/edit", data={"csrf_token": token, "name": " "})
    assert "Unnamed tenant" in client.get(f"/leases/{lid}").get_data(as_text=True)


def test_scheduled_rent_change_can_be_changed_or_cancelled(app, client):
    token = csrf(client)
    lid = q(app, "SELECT id FROM leases WHERE status = 'active' AND rent_due_day = 1 LIMIT 1")[0]
    rent = q(app, f"SELECT current_rent_cents FROM v_lease_current_rent WHERE lease_id = {lid}")[0]
    changes = lambda: q(app, f"SELECT rent_cents FROM lease_rent_changes WHERE lease_id = {lid}")
    edit = f"/leases/{lid}/edit"
    client.post(edit, data={"csrf_token": token, "rent": str(rent // 100 + 1000)})
    assert changes() == [rent + 100000]
    assert f'value="{(rent + 100000) // 100}.00"' in client.get(edit).get_data(as_text=True)  # the next bill's rent
    assert "from Oct 1, 2026" in client.get(f"/leases/{lid}").get_data(as_text=True)
    client.post(edit, data={"csrf_token": token, "rent": str(rent // 100 + 1000)})  # saving as is changes nothing
    assert changes() == [rent + 100000]
    r = client.post(edit, data={"csrf_token": token, "rent": str(rent // 100)}, follow_redirects=True)
    assert changes() == [] and "cancelled" in r.get_data(as_text=True)


def test_tenant_without_property_can_be_linked_later(app, client):
    token = csrf(client)
    r = client.post("/tenants/new", data={"csrf_token": token, "name": "Sam Lee"})
    tid = int(r.headers["Location"].rstrip("/").split("/")[-1])
    page = client.get(f"/tenants/{tid}").get_data(as_text=True)
    assert "Not linked to a unit" in page
    assert "No unit" in client.get("/tenants").get_data(as_text=True)
    pid = client.post("/properties/new", data={"csrf_token": token, "name": "9 Oak Ct"}).headers["Location"].split("/")[-1]
    unit = q(app, f"SELECT id FROM units WHERE property_id = {pid}")[0]
    r = client.post("/tenants/new", data={"csrf_token": token, "tenant_id": tid, "unit_id": unit, "rent": "900"})
    assert "/leases/" in r.headers["Location"]
    assert client.get(f"/tenants/{tid}").status_code == 302  # now goes to their account page


def test_units_not_properties(app, client):
    token = csrf(client)
    assert q(app, "SELECT MAX(n) FROM (SELECT COUNT(*) AS n FROM units GROUP BY property_id)") == [1]
    home = client.get("/").get_data(as_text=True)
    assert ">Units<" in home and ">Properties<" not in home and "Add unit" in home
    listing = client.get("/properties").get_data(as_text=True)
    assert "<h1>Units</h1>" in listing and "Add tenant" in listing and "Record payment" in listing
    form = client.get("/properties/new").get_data(as_text=True)
    assert 'name="units"' not in form and "Add unit" in form
    r = client.post("/properties/new", data={"csrf_token": token, "name": "Unit 2B Sunrise Apartments"})
    pid = int(r.headers["Location"].rstrip("/").split("/")[-1])
    page = client.get(f"/properties/{pid}").get_data(as_text=True)
    assert "Add a unit" not in page and "No tenant yet" in page and "Add tenant" in page
    assert client.post(f"/properties/{pid}/units", data={"csrf_token": token}).status_code in (404, 405)
    r = client.post(f"/delete/property/{pid}", data={"csrf_token": token}, follow_redirects=True)
    assert "Deleted Unit 2B Sunrise Apartments" in r.get_data(as_text=True)
    assert q(app, f"SELECT COUNT(*) FROM properties WHERE id = {pid}") == [0]


def test_old_multi_unit_properties_are_split_on_startup(app):
    from rental_tracker.services import portfolio
    conn = dbmod.connect(app.extensions["rental_tracker"].data.db)
    pid = portfolio.save_property(conn, None, {"name": "Old Building"}, unit_labels=["A", "B"])
    conn.close()
    app.extensions["rental_tracker"].last_catch_up = None  # as if the app just started
    c = app.test_client()
    c.get(f"/auth?token={TOKEN}")
    page = c.get("/properties?q=Old").get_data(as_text=True)
    assert "Old Building · A" in page and "Old Building · B" in page
    assert q(app, f"SELECT COUNT(*) FROM units WHERE property_id = {pid}") == [1]


def test_late_page_and_fees(app, client):
    token = csrf(client)
    page = client.get("/late").get_data(as_text=True)
    assert "overdue" in page
    keys = re.findall(r'name="key" value="([^"]+)"', page)
    assert keys, "demo tenants have late fees"
    client.post("/late", data={"csrf_token": token, "action": "waive", "key": keys[:1]})
    client.post("/late", data={"csrf_token": token, "action": "approve", "key": keys[1:2]})
    page = client.get("/late").get_data(as_text=True)
    assert keys[0] not in page and keys[1] not in page


def test_settings(client):
    token = csrf(client)
    client.post("/settings", data={"csrf_token": token, "business_name": "Sakai Rentals", "default_late_fee_cents": "50",
                                   "default_grace_days": "3", "late_fee_mode": "review"})
    page = client.get("/settings").get_data(as_text=True)
    assert "Sakai Rentals" in page and 'value="50.00"' in page
    assert 'value="50.00"' in client.get("/tenants/new").get_data(as_text=True)  # used for new tenants


def test_quit_button(app, client):
    token = csrf(client)
    assert client.post("/quit", data={"csrf_token": token}).status_code == 404  # not started by the launcher
    called = []
    app.extensions["rental_tracker"].on_quit = lambda: called.append(True)
    assert "Quit" in client.get("/").get_data(as_text=True)
    resp = client.post("/quit", data={"csrf_token": token})
    assert resp.status_code == 200 and "has closed" in resp.get_data(as_text=True)
    time.sleep(0.5)
    assert called == [True]


def test_no_back_button_or_vacant_units(app):
    c = app.test_client()
    c.get(f"/auth?token={TOKEN}&window=1")
    assert "data-back" not in c.get("/").get_data(as_text=True)
    assert "← Back" not in c.get(f"/payments/{q(app, 'SELECT id FROM payments LIMIT 1')[0]}/receipt").get_data(as_text=True)
    props = c.get("/properties").get_data(as_text=True)
    assert 'name="vacant"' not in props and "vacant" not in props.lower()
    home = c.get("/").get_data(as_text=True)
    assert "Vacant" not in home and "Occupied" not in home
    assert "vacant" not in c.get("/reports/rent-roll").get_data(as_text=True).lower()
    assert "Vacant units" not in c.get("/reports").get_data(as_text=True)


def test_payment_methods_and_no_ref(app, client):
    token = csrf(client)
    lid = q(app, "SELECT id FROM leases WHERE status = 'active' LIMIT 1")[0]
    page = client.get(f"/leases/{lid}").get_data(as_text=True)
    methods = re.findall(r'<option value="([a-z_]+)"', page.split('name="method"')[1].split("</select>")[0])
    assert methods == ["cash", "check", "bank_transfer", "gcash", "other"]
    assert 'name="reference"' not in page and 'name="method_other"' in page
    client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": "100", "method": "gcash"})
    client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": "100", "method": "other",
                                                "method_other": "  Maya   wallet "})
    client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": "100", "method": "cash",
                                                "method_other": "ignored"})
    rows = [tuple(r) for r in dbmod.connect(app.extensions["rental_tracker"].data.db).execute(
        f"SELECT method, method_other FROM payments WHERE lease_id = {lid} ORDER BY id DESC LIMIT 3")]
    assert rows == [("cash", None), ("other", "Maya wallet"), ("gcash", None)]
    other = q(app, f"SELECT id FROM payments WHERE lease_id = {lid} AND method = 'other' ORDER BY id DESC LIMIT 1")[0]
    assert "Maya wallet" in client.get(f"/payments/{other}/receipt").get_data(as_text=True)
    assert "Payment — Maya wallet" in client.get(f"/leases/{lid}").get_data(as_text=True)
    listing = client.get("/payments").get_data(as_text=True)
    assert "Maya wallet" in listing and "GCash" in listing and "<th>Ref</th>" not in listing
    assert "Maya wallet" in [c.value for c in sheet(client.get("/payments?format=xlsx"))["E"]]
    grid = client.get("/rent-day").get_data(as_text=True)
    assert "<th>Ref</th>" not in grid and 'name="reference"' not in grid
    r = client.post("/rent-day/pay", data={"lease_id": lid, "amount": "50", "method": "other", "method_other": "Coins.ph",
                                           "period": "2026-09"}, headers={"HX-Request": "true", "X-CSRF-Token": token})
    assert "Saved" in r.get_data(as_text=True)
    assert q(app, f"SELECT method_other FROM payments WHERE lease_id = {lid} ORDER BY id DESC LIMIT 1") == ["Coins.ph"]
    assert 'value="Coins.ph"' in r.get_data(as_text=True)  # the row keeps it for next time
    client.post(f"/leases/{lid}/payment", data={"csrf_token": token, "amount": "1", "method": "money_order"})
    assert q(app, f"SELECT COUNT(*) FROM payments WHERE lease_id = {lid} AND method = 'money_order'")[0] == 0


def test_money_is_in_pesos(client):
    for page in ("/", "/properties", "/tenants", "/rent-day", "/payments", "/late", "/reports/rent-roll"):
        html = client.get(page).get_data(as_text=True)
        assert "₱" in html, page
        assert not re.search(r"\$\d", html), page


def test_deletes_happen_without_asking(app, client):
    token = csrf(client)
    lease_id = q(app, "SELECT id FROM leases WHERE status = 'active' LIMIT 1")[0]
    page = client.get(f"/leases/{lease_id}").get_data(as_text=True)
    assert "delete-payment" in page and "delete-charge" in page and f"/delete/lease/{lease_id}" in page
    pid = q(app, "SELECT id FROM properties LIMIT 1")[0]
    pages = page + client.get("/payments").get_data(as_text=True) + client.get(f"/properties/{pid}").get_data(as_text=True)
    for form in re.findall(r"<form[^>]*>", pages):  # no "are you sure?" on any delete
        if "delete" in form or "remove" in form:
            assert "data-confirm" not in form, form
    assert client.get(f"/delete/lease/{lease_id}").status_code == 405  # no confirmation page any more
    assert client.post("/delete/nonsense/1", data={"csrf_token": token}).status_code == 404
    r = client.post(f"/delete/lease/{lease_id}", data={"csrf_token": token})
    assert r.status_code == 302 and "/properties/" in r.headers["Location"]
    assert q(app, f"SELECT COUNT(*) FROM charges WHERE lease_id = {lease_id}")[0] == 0
    backups = list((app.extensions["rental_tracker"].data.backups / "snapshots").glob("*before-delete-lease*"))
    assert backups
    pid = q(app, "SELECT id FROM properties LIMIT 1")[0]
    client.post(f"/delete/property/{pid}", data={"csrf_token": token})
    assert q(app, f"SELECT COUNT(*) FROM properties WHERE id = {pid}")[0] == 0


def test_delete_single_entries(app, client):
    token = csrf(client)
    pay = q(app, "SELECT id FROM payments LIMIT 1")[0]
    lease = q(app, f"SELECT lease_id FROM payments WHERE id = {pay}")[0]
    other = q(app, f"SELECT id FROM leases WHERE id <> {lease} LIMIT 1")[0]
    assert client.post(f"/leases/{other}/delete-payment/{pay}", data={"csrf_token": token}).status_code == 404
    client.post(f"/leases/{lease}/delete-payment/{pay}", data={"csrf_token": token})
    assert q(app, f"SELECT COUNT(*) FROM payments WHERE id = {pay}")[0] == 0
    pay2 = q(app, "SELECT id FROM payments LIMIT 1")[0]
    r = client.post(f"/payments/{pay2}/delete", data={"csrf_token": token, "next": "/payments?start=2026-01-01"})
    assert r.headers["Location"].endswith("/payments?start=2026-01-01")
    assert q(app, "SELECT COUNT(*) FROM audit_log WHERE action = 'delete'")[0] >= 2


def test_security_headers(client):
    resp = client.get("/")
    assert "script-src 'self'" in resp.headers["Content-Security-Policy"]
    assert resp.headers["X-Frame-Options"] == "DENY"


def sheet(resp, name=None):
    """Open an Excel download from the app."""
    from openpyxl import load_workbook
    assert resp.status_code == 200 and resp.mimetype.endswith("spreadsheetml.sheet"), resp.status_code
    disposition = resp.headers["Content-Disposition"]
    assert disposition.startswith('attachment; filename="') and disposition.endswith('.xlsx"'), disposition
    wb = load_workbook(io.BytesIO(resp.data))
    return wb[name] if name else wb.active


def test_excel_exports(app, client):
    from datetime import datetime
    from decimal import Decimal
    from openpyxl import load_workbook
    token = csrf(client)
    # tenants: real numbers shown as pesos, real dates, a header row
    ws = sheet(client.get("/tenants?format=xlsx"))
    assert [c.value for c in ws[1]] == ["Tenant", "Unit", "Phone", "Rent", "Balance", "Overdue", "Days late",
                                        "Last paid", "Status"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    rent = ws["D2"]
    assert isinstance(rent.value, (int, float, Decimal)) and "₱" in rent.number_format
    assert any(isinstance(c.value, datetime) for c in ws["H"][1:])
    assert ws.max_row - 1 == len(client.get("/tenants").get_data(as_text=True).split("<tbody>")[1].split("</tbody>")[0]
                                 .split("<tr>")) - 1
    # payments: a total row that adds up
    ws = sheet(client.get("/payments?format=xlsx&start=2026-01-01&end=2026-12-31"))
    amounts = [c.value for c in ws["F"][1:-1]]
    assert ws.cell(ws.max_row, 1).value == "Total" and ws.cell(ws.max_row, 6).value == sum(amounts)
    # reports keep their columns and totals
    ws = sheet(client.get("/reports/rent-roll?format=xlsx"))
    assert ws["A1"].value == "Unit" and ws.cell(ws.max_row, 1).value == "Total"
    # one tenant's payment history, oldest first, with the running balance
    lid = q(app, "SELECT id FROM leases WHERE status = 'active' LIMIT 1")[0]
    assert "Export to Excel" in client.get(f"/leases/{lid}").get_data(as_text=True)
    ws = sheet(client.get(f"/leases/{lid}/statement?format=xlsx"))
    assert [c.value for c in ws[1]] == ["Date", "What", "Charged", "Paid", "Balance"]
    balance = q(app, f"SELECT balance_cents FROM v_lease_balances WHERE lease_id = {lid}")[0]
    assert ws.cell(ws.max_row, 5).value * 100 == balance
    # everything in one workbook
    wb = load_workbook(io.BytesIO(client.get("/reports/everything.xlsx").data))
    assert wb.sheetnames == ["Units", "Tenants", "Payments", "History"]
    assert wb["Units"].max_row - 1 == q(app, "SELECT COUNT(*) FROM properties")[0]
    assert wb["Payments"].max_row - 2 == q(app, "SELECT COUNT(*) FROM payments")[0]
    assert "Export everything to Excel" in client.get("/reports").get_data(as_text=True)
    # text is never turned into a formula
    client.post(f"/leases/{lid}/edit", data={"csrf_token": token, "name": "=1+2"})
    ws = sheet(client.get("/tenants?format=xlsx"))
    cell = next(c for c in ws["A"] if str(c.value).startswith("=1+2"))
    assert cell.data_type == "s"
