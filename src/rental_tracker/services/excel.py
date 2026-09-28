"""Excel (.xlsx) downloads for lists, reports, a tenant's history, or everything at once.

Amounts are real numbers shown as pesos and dates are real dates, so Excel can sort,
filter and add them up. Text is always stored as text: a name starting with "=" can't
become a formula.
"""
from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = '"₱"#,##0.00;-"₱"#,##0.00'
DATE_FORMAT = "mmm d, yyyy"
PCT_FORMAT = '0.0"%"'  # values are already percentages, e.g. 91.5
FORMATS = {"money": MONEY_FORMAT, "date": DATE_FORMAT, "pct": PCT_FORMAT, "int": "0"}
_BAD_TITLE_CHARS = str.maketrans({c: " " for c in "[]:*?/\\"})


@dataclass
class Sheet:
    """One worksheet: a header row, data rows and an optional totals row.

    columns are (heading, kind) with kind text | money (cents) | date (ISO text) | pct | int.
    """
    title: str
    columns: list[tuple[str, str]]
    rows: list[list[Any]]
    totals: list[Any] | None = None


def _value(value: Any, kind: str) -> Any:
    if value is None or value == "":
        return None
    if kind == "money":
        return Decimal(int(value)) / 100
    if kind == "date":
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return str(value)
    if kind == "pct":
        return float(value)
    if kind == "int":
        return int(value)
    return str(value)


def _title(text: str, used: set[str]) -> str:
    """Excel sheet names: at most 31 characters, none of []:*?/\\, and unique."""
    base = " ".join(str(text).translate(_BAD_TITLE_CHARS).split())[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def workbook(*sheets: Sheet) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EDF4")
    for sheet in sheets:
        ws = wb.create_sheet(_title(sheet.title, used))
        kinds = [kind for _, kind in sheet.columns]
        ws.append([heading for heading, _ in sheet.columns])
        for cell in ws[1]:
            cell.font, cell.fill = bold, header_fill
        body = sheet.rows + ([sheet.totals] if sheet.totals else [])
        widths = [len(heading) for heading, _ in sheet.columns]
        for values in body:
            ws.append([_value(v, k) for v, k in zip(values, kinds)])
            for i, (cell, kind) in enumerate(zip(ws[ws.max_row], kinds)):
                if cell.value is None:
                    continue
                if kind in FORMATS:
                    cell.number_format = FORMATS[kind]
                    shown = 14 if kind in ("money", "date") else 8
                else:
                    cell.data_type = "s"  # text stays text, even if it starts with "="
                    shown = len(str(cell.value))
                widths[i] = max(widths[i], shown)
        if sheet.totals:
            for cell in ws[ws.max_row]:
                cell.font = bold
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)
        ws.freeze_panes = "A2"
        last_data_row = ws.max_row - (1 if sheet.totals else 0)
        if sheet.columns and last_data_row > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{last_data_row}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- the sheets the app offers ---------------------------------------------------------------

STATE_LABELS = {"late": "Late", "owes": "Owes", "paid": "Paid up", "credit": "Credit", "none": "No unit"}


def tenants_sheet(rows: list[dict]) -> Sheet:
    return Sheet("Tenants", [("Tenant", "text"), ("Unit", "text"), ("Phone", "text"), ("Rent", "money"),
                             ("Balance", "money"), ("Overdue", "money"), ("Days late", "int"),
                             ("Last paid", "date"), ("Status", "text")],
                 [[r["names"], r["property_code"], r["phone"], r["current_rent_cents"], r["balance_cents"],
                   r["past_due_cents"], r["days_late"] or None, r["last_paid_on"],
                   STATE_LABELS.get(r["state"], r["state"])] for r in rows])


def payments_sheet(rows: list, method_name) -> Sheet:
    columns = [("Date", "date"), ("Receipt", "text"), ("Unit", "text"), ("Tenant", "text"), ("Method", "text"),
               ("Amount", "money")]
    voided = any(r["voided_at"] for r in rows)  # only payments voided by older versions
    if voided:
        columns.append(("Voided", "text"))
    out = []
    for r in rows:
        line = [r["received_date"], r["receipt_number"], r["property_code"], r["tenants"],
                method_name(r["method"], r["method_other"]), r["amount_cents"]]
        if voided:
            line.append(r["void_reason"] if r["voided_at"] else None)
        out.append(line)
    total = sum(r["amount_cents"] for r in rows if not r["voided_at"])
    return Sheet("Payments", columns, out, ["Total", None, None, None, None, total] + ([None] if voided else []))


def history_sheet(title: str, entries: list[dict], start: str = "", opening: int = 0) -> Sheet:
    """A tenant's payment history, oldest first, with the running balance.

    From a start date, the first line is the balance carried forward from before it.
    """
    rows = [[start, "Balance forward", None, None, opening]] if start else []
    rows += [[e["date"], e["description"], e["charge"] or None, e["credit"] or None, e["balance"]]
             for e in entries]
    return Sheet(title, [("Date", "date"), ("What", "text"), ("Charged", "money"), ("Paid", "money"),
                         ("Balance", "money")], rows)


def everything(conn: sqlite3.Connection, today: date) -> bytes:
    """One workbook with every unit, tenant, payment and history line."""
    from . import ledger, portfolio, tenants
    rows = tenants.tenancies(conn, today=today, status="all")
    by_unit: dict[int, list[dict]] = {}
    for r in rows:
        if r["property_id"] and r["status"] in ("active", "month_to_month", "future"):
            by_unit.setdefault(r["property_id"], []).append(r)
    units = Sheet("Units", [("Unit", "text"), ("City", "text"), ("Province", "text"), ("ZIP code", "text"),
                            ("Tenant", "text"), ("Rent", "money"), ("Owed", "money"), ("Notes", "text")],
                  [[p["name"], p["city"], p["state"], p["postal_code"],
                    ", ".join(t["names"] or "" for t in by_unit.get(p["id"], [])), p["rent_cents"],
                    p["balance_cents"], p["notes"]]
                   for p in portfolio.list_properties(conn, status="all", sort="name")])
    history = []
    for r in rows:
        if not r["lease_id"]:
            continue
        for e in ledger.ledger_entries(conn, r["lease_id"]):
            history.append([e["date"], r["property_code"], r["names"], e["description"],
                            e["charge"] or None, e["credit"] or None])
    history.sort(key=lambda line: (line[0], line[1] or "", line[2] or ""))
    return workbook(
        units, tenants_sheet(rows),
        payments_sheet(ledger.list_payments(conn, limit=10 ** 9), ledger.method_name),
        Sheet("History", [("Date", "date"), ("Unit", "text"), ("Tenant", "text"), ("What", "text"),
                          ("Charged", "money"), ("Paid", "money")], history))
